import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = "docs/PROJECT_REPORT_2PAGE.md"
OUT = sys.argv[1] if len(sys.argv) > 1 else "docs/ClearCut_Report_2pages.docx"
BODY = float(sys.argv[2]) if len(sys.argv) > 2 else 9.5
FONT = "Calibri"
TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`)")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin, sec.bottom_margin = Cm(1.4), Cm(1.3)
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(BODY)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.0


def runs(p, text, size=None, bold=False, color=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        code = part.startswith("`") and part.endswith("`") and len(part) > 2
        strong = part.startswith("**") and part.endswith("**") and len(part) > 4
        if code:
            part = part[1:-1]
        elif strong:
            part = part[2:-2]
        run = p.add_run(part)
        run.bold = bold or strong
        if code:
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if size:
            run.font.size = Pt(size - 0.5 if code else size)
        if color:
            run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def flush_table(rows):
    rows = [r for r in rows if not re.fullmatch(r"\|?[\s:\-|]+\|?", r)]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    table = doc.add_table(rows=len(cells), cols=len(cells[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [4.2, 13.2] if len(cells[0]) == 2 and cells[0][0] == "File" else [12.4, 5.0]
    for r, row in enumerate(cells):
        trPr = table.rows[r]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.width = Cm(widths[c])
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if r < len(cells) - 1:
                p.paragraph_format.keep_with_next = True
            runs(p, value, size=BODY - 0.5, bold=(r == 0))
            if r == 0:
                shade(cell, "DCE6F1")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


lines = open(SRC, encoding="utf-8").read().splitlines()
pending = []
for line in lines + [""]:
    if line.startswith("|"):
        pending.append(line)
        continue
    if pending:
        flush_table(pending)
        pending = []
    if not line.strip():
        continue
    if line.startswith("# "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        runs(p, line[2:], size=15, bold=True, color=(0x1F, 0x3A, 0x5F))
    elif line.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        runs(p, line[3:], size=11.5, bold=True, color=(0x1F, 0x3A, 0x5F))
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "1F3A5F")):
            bottom.set(qn(key), value)
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)
    elif re.match(r"^\s+- ", line):
        p = doc.add_paragraph(style="List Bullet 2")
        p.paragraph_format.space_after = Pt(1)
        runs(p, line.strip()[2:], size=BODY)
    elif line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        runs(p, line[2:], size=BODY)
    elif re.match(r"^\d+\. ", line):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(1)
        runs(p, re.sub(r"^\d+\. ", "", line), size=BODY)
    else:
        p = doc.add_paragraph()
        if line.startswith("**Code:**"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        runs(p, line, size=BODY)

doc.save(OUT)
print("saved", OUT)
