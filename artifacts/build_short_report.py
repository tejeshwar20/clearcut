"""Build the 2-page project report (docs/ClearCut_Report_2pages.docx)."""
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8")
OUT = "docs/ClearCut_Report_2pages.docx"
FONT = "Calibri"
BODY = 10

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21.0), Cm(29.7)
section.left_margin = section.right_margin = Cm(1.9)
section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.4)

style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(BODY)
style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
style.paragraph_format.space_after = Pt(3)
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.line_spacing = 1.0


def para(text="", bold=False, italic=False, size=None, align=None, after=None, before=None, color=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    add_runs(p, text, bold, italic, size, color)
    return p


import re

TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)")


def add_runs(p, text, bold=False, italic=False, size=None, color=None):
    """Supports **bold**, *italic* and `code` segments inside text."""
    for part in TOKEN.split(text):
        if not part:
            continue
        is_bold, is_italic, is_code = False, False, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, is_bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, is_code = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, is_italic = part[1:-1], True
        run = p.add_run(part)
        run.bold = bold or is_bold
        run.italic = italic or is_italic
        if is_code:
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if size:
            run.font.size = Pt(size - 0.5 if is_code else size)
        if color:
            run.font.color.rgb = RGBColor(*color)


def heading(text):
    p = para(text, bold=True, size=11.5, before=5, after=2, color=(0x1F, 0x3A, 0x5F))
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "1F3A5F")):
        bottom.set(qn(key), value)
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.55)
    add_runs(p, text)
    for run in p.runs:
        run.font.size = Pt(BODY)
    return p


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def table(rows, widths, header=True, size=9):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = t.cell(r, c)
            cell.width = Cm(widths[c])
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, value, bold=(header and r == 0), size=size)
            if header and r == 0:
                shade(cell, "DCE6F1")
    for r, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        if r < len(t.rows) - 1:
            for cell in row.cells:
                for cp in cell.paragraphs:
                    cp.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


para("ClearCut: Rule-Based Spoiler-Free Rewriting of Movie and Series Reviews", bold=True, size=15,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=1, color=(0x1F, 0x3A, 0x5F))
para("Natural Language Processing Project Report", italic=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
para("Tejeshwar  |  [Roll number, programme, institution]  |  Submitted to: [Professor, department]  |  23 October 2026",
     size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
para("Code: https://github.com/tejeshwar20/clearcut", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=3)

heading("Abstract")
para("Reviews often reveal plot twists, deaths and endings. ClearCut is a web application that takes a pasted English review "
     "and returns the same review with only the spoiler sentences rewritten; every other sentence is kept word for word. "
     "It is a rule-based NLP pipeline: spaCy dependency parsing decides how much of a sentence to remove, curated pattern sets "
     "detect plot disclosures, and a rewriting step keeps the reviewer's reaction while replacing the revealed fact with a neutral "
     "phrase. A final check withholds the output if any disclosure remains. On developer-written tests the system leaked 0 of 25 "
     "spoilers, changed 0 of 20 harmless sentences, kept 48 of 48 harmless sentences around a spoiler, and responded in about 20 ms "
     "(short) to 130 ms (300 words). These tests are not an independent benchmark.")

heading("1. Introduction")
para("Readers who want a reviewer's opinion but not the plot have few options. Sites offer spoiler tags, and extensions blur "
     "keyword matches, but both hide the opinion along with the plot. ClearCut instead rewrites the spoiler and keeps the "
     "opinion. For example, \u201cGreat visuals. It was all a simulation the whole time, which blew my mind.\u201d becomes "
     "\u201cGreat visuals. The big surprise blew my mind.\u201d The reviewer's reaction survives and the revealed fact does not.")

heading("2. Problem statement")
para("Given a free-text review, produce a version that **(a)** discloses no plot information, **(b)** keeps every sentence that "
     "does not disclose plot exactly as written, and **(c)** keeps the reviewer's reaction wherever this is safe. "
     "Existing spoiler-detection research mostly labels text as spoiler / not spoiler; this project goes further and "
     "rewrites the text, under the strict constraint that harmless text is never altered.")

heading("3. Solution approach")
para("**Model.** There is no trained model. A transparent rule-based pipeline was chosen because no labelled spoiler data was "
     "available, because \u201cother sentences stay identical\u201d is easy to guarantee with rules but not with a generative "
     "model, and because every decision can be traced to a named pattern. It runs offline and needs no API key.", after=2)
flow = ["Pasted review", "Parse (spaCy)", "Detect spoiler cues", "Keep / cut / rewrite",
        "Rewrite reaction", "Safety check", "Spoiler-free review"]
ft = doc.add_table(rows=1, cols=len(flow))
ft.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, label in enumerate(flow):
    cell = ft.cell(0, i)
    cell.text = ""
    fp = cell.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(0)
    run = fp.add_run(label + ("" if i == len(flow) - 1 else "  \u2192"))
    run.font.size = Pt(7.5)
    run.bold = True
    shade(cell, "E8F0E4" if i in (0, len(flow) - 1) else "DCE6F1")
doc.add_paragraph().paragraph_format.space_after = Pt(0)
para("**Figure 1.** Processing pipeline. Sentences with no spoiler cue skip straight through unchanged.", size=8.5, italic=True, after=3)
para("**Pipeline** (spaCy 3.8 `en_core_web_sm`, regular expressions, VADER 3.3.2, Flask/Waitress):", after=1)
bullet("**Parse:** split into sentences; obtain part-of-speech tags, lemmas and the dependency tree.")
bullet("**Detect:** match five cue sets \u2013 events (deaths, betrayals), reveals (\u201cturns out\u201d, \u201cthe climax\u201d), hidden "
       "reveals (identity and family reveals, \u201call along\u201d, dream or simulation twists, \u201cSpoiler:\u201d), outcomes, and "
       "character fates. Genre labels such as \u201crevenge thriller\u201d are exempt.")
bullet("**Scope:** no cue \u2192 sentence kept unchanged. Cue in a subordinate clause \u2192 only that clause is removed. "
       "Cue in the main clause or a severe reveal \u2192 the sentence is rewritten.")
bullet("**Rewrite:** keep the reviewer's reaction and point it at a neutral phrase (\u201cthe big surprise\u201d); if no opinion can be "
       "kept, use a neutral line chosen by spoiler type (ending, climax, twist, fate). Lines are never repeated in one review, "
       "and consecutive spoilers are merged into one line.")
bullet("**Safety check:** the output is re-scanned for disclosure words and long copied runs of the source; if any are found, "
       "nothing is shown (the system fails closed).")
para("**Design decisions.** A sentence-embedding \u201csemantic\u201d layer was built and then removed because it altered harmless "
     "sentences and raised memory from about 238 MB to 398 MB. An estimated review score was also removed, because it was only a "
     "guess from wording and not part of the task.", before=2)

heading("4. Experimental results")
para("A script (`artifacts/measure.py`) runs fixed test sets through the live pipeline. Each spoiler case is a harmless opening "
     "sentence followed by a spoiler; the test lists terms that must not appear in the output. Cases were written by the "
     "developer, so the results show the rules behave as designed, not accuracy on unseen reviews.", after=3)
cap1 = para("**Table 1.** Measured results.", size=9, italic=True, after=2)
cap1.paragraph_format.keep_with_next = True
table([
    ["Measure", "Result"],
    ["Spoiler sentences tested / leaked into output", "25 / **0**"],
    ["Harmless sentences tested / changed or blocked", "20 / **0**"],
    ["Mixed reviews (harmless + spoiler + harmless): harmless sentences kept intact", "**48 of 48**"],
    ["Whole review wrongly blocked", "0"],
    ["Latency, short review (about 2 sentences): median / 95th percentile", "20 ms / 26 ms"],
    ["Latency, 297-word review: median / maximum", "127 ms / 170 ms"],
    ["Automated tests passing", "212 of 212"],
    ["Earlier aspect-sentiment component (`evaluate.py`, 26 synthetic examples)", "spoiler accuracy 0.92; sentiment accuracy 0.89"],
], widths=[12.2, 5.0])
cap2 = para("**Table 2.** Worked examples (actual outputs of the final system).", size=9, italic=True, after=2)
cap2.paragraph_format.keep_with_next = True
table([
    ["Input", "Output"],
    ["The acting is superb. Darth Vader is Luke's father, and that scene still gives me chills. The pacing is tight.",
     "The acting is superb. The big surprise still gives me chills. The pacing is tight."],
    ["Fun ride. Turns out the mentor was the villain, which I didn't see coming.",
     "Fun ride. I didn't see the big surprise coming."],
    ["The cast is superb and the songs are catchy. In the end the hero dies saving his village. The cinematography is gorgeous.",
     "The cast is superb and the songs are catchy. The ending is best experienced without knowing what's coming. The cinematography is gorgeous."],
    ["The visual effects are stunning, but I hated that the main character dies at the end.",
     "The visual effects are stunning."],
    ["A tense, gripping war drama with strong performances.", "(unchanged: no spoiler)"],
], widths=[8.6, 8.6], size=8.5)
para("**Finding.** The first run of the 25-case set exposed one leak (\u201cAmy faked her own disappearance to frame Nick\u201d), "
     "although an earlier 16-case probe had shown none. The patterns were extended, a test was added, and the re-run gave 0 leaks. "
     "A real failure was also fixed: a pasted plot summary of *Devara: Part 1* was blocked entirely because a rewrite produced "
     "\u201cThe big surprise is dead\u201d; it now keeps the setup verbatim and replaces only the twist and ending.", after=2)
para("**Limitations.** Rules cover only anticipated phrasings, so an unusual spoiler can pass; replacement lines are generic by "
     "design; only English is supported; and accuracy on real reviews is unknown until tested on an independent, labelled set "
     "(planned: 100 or more reviews, two annotators).")

heading("5. Data source")
para("The system uses **no training data**. Test and evaluation data are as follows.", after=1)
bullet("**Test reviews and spoiler sentences:** written by the developer, based on the plots of well-known films (for example "
       "*The Sixth Sense*, *Fight Club*, *Parasite*, *Gone Girl*) and on reviews that failed during use (*Jailer*, *Irumudi*, *Devara*). "
       "They live in `tests/` and `artifacts/measure.py`.")
bullet("**Evaluation file:** `data/evaluation_examples.json` holds 26 short synthetic sentences used by `evaluate.py`; it is a sanity check, not a benchmark.")
bullet("**Libraries and lexicons:** spaCy `en_core_web_sm` 3.8.0 (parser) and the VADER sentiment lexicon (published rule-based model, "
       "Hutto and Gilbert, 2014), both used unmodified except for a small added review vocabulary.")
bullet("**User input:** pasted reviews are processed in memory and are not stored by the pasted-review path.")
para("Source code, tests and the full 14-section report: https://github.com/tejeshwar20/clearcut", size=9, italic=True, before=2)

doc.save(OUT)
print("saved", OUT)
